import os
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd
from pandas import ExcelWriter
from pandas.io.parsers import read_csv
from xlsxwriter import *
import openpyxl as opxl



def createMergedMultipleCampsFile(FileList, termcode, nextdir2):
    print('\nThe list has been passed in.  Press enter to continue: ')
    input()

    cwd = os.getcwd()  # Current Working Dirctory
    #nextdir =  termcode
    #nextdir2 = input('Enter the month of the folder in which you want the merged file to be saved, (usually the current month):  ').capitalize()
    path = os.path.join(cwd, termcode, nextdir2)

    df_list = []
    for f in FileList:
        f = os.path.join(path, f)
        df = read_csv(f)
        df_list.append(df)

    mergedFile = pd.concat(df_list)

    mergedFile = pd.DataFrame(data=mergedFile)

    comments = {'Course/reason': ['Not Required course', 'NSCC 1010', 'Lower level math that is prereq for MATH 1910 sequence', 'Lower level math that is prereq for MATH sequence', 'Lower level math that is prereq for Math sequence, 2nd degree', 'Lower level math that is prereq for Science sequence', 'Repeat, 2nd attempt at class', 'Repeat replacing transfer class', 'repeat that is beyond 2nd attempt (change number to match)', 'Course that has been routinely subbed by dean, email dean', 'GPA below 2.0 and being used to boost', 'Required for residency', 'When an exception has been added to force class into correct slot (ex. Health Sciences UNPA)', 'CPoS pulled class to report in error-possible change after report ran', 'Required for Degree but student dropped it', 'Class is required change to Banner or DW needs to be made by Records', 'Class is required, Change made to Banner for 2nd Degree', '1st 7 week, student failed', '1st 7 week, was required', 'Accelerated course; student earned failing grade'], 'Comment': [
        'Not required for degree', 'Is required; FYE added to SGASADD & DW updated', 'Is required; hidden prereq for Math 1910 sequence', 'Is required; hidden prereq for Math sequence', 'Is required; hidden prereq for Math sequence in 2nd degree', 'Is required; hidden prereq for required Science sequence', 'Is required; Spring 2022 is 2nd attempt', 'Is required; will replace transfer grade', 'Is required; Spring 2022 is 5th attempt; can\'t pay', 'Required for Degree, Substituiton incoming from Dean', 'Is Required; Repeat being used to boost GPA to minimum 2.0', 'Is required, needed to meet residency', 'Is required, DW has been updated', 'Is required; DW is showing correctly', 'Is required, student dropped class', 'Is required, DW is being corrected', 'Is required, DW is showing correctly, 2ND MAJOR CORRECTED', '1st 7 week, was required, student failed course', '1st 7 weeks, was required for degree', 'Required; 7 week course-student failed']}
    comments = pd.DataFrame(data=comments)

    # Get the columns from the DataFrame that are wanted on the report
    mergedFile = mergedFile.iloc[:, :]

    cols = list(mergedFile.columns.values)
    mergedFile = mergedFile[[cols[0]]+[cols[1]]+[cols[2]]+[cols[3]]+[cols[4]]+[cols[5]]+[cols[6]]+[
        cols[7]]+[cols[10]]+[cols[12]]+[cols[13]]+[cols[14]]+[cols[15]]+[cols[18]]+[cols[47]]]

    mergedFile = mergedFile.rename(
        columns={
            'stu_id': "A Number",
            'stu_sort_name': "Name",
            'cpos_term': "Term",
            'course': "Course",
            'crn': "CRN",
            'repeat_code': "Repeat?",
            'repeat_count': "Repeat Count",
            'stu_program': "Program",
            'concentration': "Concentration",
            'veteran': "Veteran?",
            'repeat_count_no_wd': 'Repeat Count No WD',
            'stvcoll_desc': 'STVCOLL Desc',
            'stu_camp_email': 'Campus eMail',
            'Activity_Date': 'Activity Date',
            'Max_Registration_Date': 'Max Reg Date'
        }
    )
    mergedFile['Comment'] = '  '

    # Remove duplicate records (exculing activity date) from the file
    unDuplicated = mergedFile.drop_duplicates(subset=["A Number", "Name", "Term", "Course", "CRN", "Repeat?", "Repeat Count",
                                              "Program", "Concentration", "Veteran?", 'Repeat Count No WD', 'STVCOLL Desc', 'Campus eMail'], keep='last')
    input('\nThe unduplicated report has now been created.  The most recent record was kept and all previous duplicated records were removed.  Press enter to see a preview:  ')
    print(unDuplicated)
    print('\nAfter reviewing the preview, press enter to continue:  ')
    input()

    # Write DataFrame to Excel File
    file_name = input('Enter the name you wish to give the merged file:  ')
    if not file_name.endswith('.xlsx'):
        file_name = file_name+'.xlsx'
    else:
        pass

    writer = pd.ExcelWriter(file_name)
    with writer as writer:
        unDuplicated.to_excel(writer, sheet_name="Merged CPoS", index=False)
        comments.to_excel(writer, sheet_name="Comments", index=False)

    # change the dirctory to the proper location from the CPoS general directory
    source = os.path.join(cwd, file_name)
    destination = os.path.join(cwd, termcode, nextdir2, file_name)
    os.rename(source, destination)

    # Give confirmation and let the user know the file is ready to be worked
    print(
        f'\nDone!  The merged file can be found by following this path: {path}.\n')
    input('\nPress enter for an important conclucing message.  ')
    print('\nNOTE:  Be sure to delete the .CSV files you imported into this merged file. ')

    print('\nHave a great day!\n')

    exit()


def main(termcode, nextdir2):
    # Import CSV File
    infile = input(
        '\nEnter the file\'s name you wish to load (format=mmdd_cpos_termcode):  \n')

    if not infile.endswith('csv'):
        infile = infile+'.csv'
    else:
        pass
    cwd = os.getcwd()  # Current Working Dirctory
    # nextdir = infile[-10:-4] #input('Enter the registration term code for this report:  ')
    #nextdir2 = input('Enter the month of which the file was created, (usually the current month):  ').capitalize()
    infile = os.path.join(cwd, termcode, nextdir2, infile)
    
    try:
        open(infile, 'r')
    except FileNotFoundError:
        input('\nERROR:  The file was not found; press enter to start over:  ')
        main(termcode, nextdir2)
    except:
        input('\Something besides a FileNotFoundError went wrong.  Press enter to start over:  ')
        main(termcode, nextdir2)

    print(
        f'\nThe file {infile} has been imported into this program, next I will convert the file to an Excel File from CSV.')
    input('\nPress enter to continue: ')

    # Convert CSV File to Pandas DataFrame
    # The read_csv method converts a csv file to a dataframe
    wksht = pd.read_csv(infile)

    # Remove unwanted columns and add comment column
    wksht = wksht.iloc[:, :]

    cols = list(wksht.columns.values)
    wksht = wksht[[cols[0]]+[cols[1]]+[cols[2]]+[cols[3]]+[cols[4]]+[cols[5]]+[cols[6]] +
                  [cols[7]]+[cols[10]]+[cols[12]]+[cols[13]]+[cols[14]]+[cols[15]]+[cols[18]]+[cols[47]]]

    wksht = wksht.rename(
        columns={
            'stu_id': "A Number",
            'stu_sort_name': "Name",
            'cpos_term': "Term",
            'course': "Course",
            'crn': "CRN",
            'repeat_code': "Repeat?",
            'repeat_count': "Repeat Count",
            'stu_program': "Program",
            'concentration': "Concentration",
            'veteran': "Veteran?",
            'repeat_count_no_wd': 'Repeat Count No WD',
            'stvcoll_desc': 'STVCOLL Desc',
            'stu_camp_email': 'Campus eMail',
            'Activity_Date': 'Activity Date',
            'Max_Registration_Date': 'Max Reg Date'
        }
    )
    wksht['Comment'] = '  '

    input('\nThe file has been created, press enter to see a preview of the file:  ')
    print(wksht.head(10))
    input('\n After reviewing the preview press enter again to finalize the file:   ')

    comments = {'Course/reason': ['Not Required course', 'NSCC 1010', 'Lower level math that is prereq for MATH 1910 sequence', 'Lower level math that is prereq for MATH sequence', 'Lower level math that is prereq for Math sequence, 2nd degree', 'Lower level math that is prereq for Science sequence', 'Repeat, 2nd attempt at class', 'Repeat replacing transfer class', 'repeat that is beyond 2nd attempt (change number to match)', 'Course that has been routinely subbed by dean, email dean', 'GPA below 2.0 and being used to boost', 'Required for residency', 'When an exception has been added to force class into correct slot (ex. Health Sciences UNPA)', 'CPoS pulled class to report in error-possible change after report ran', 'Required for Degree but student dropped it', 'Class is required change to Banner or DW needs to be made by Records', 'Class is required, Change made to Banner for 2nd Degree', '1st 7 week, student failed', '1st 7 week, was required', 'Accelerated course; student earned failing grade'], 'Comment': [
        'Not required for degree', 'Is required; FYE added to SGASADD & DW updated', 'Is required; hidden prereq for Math 1910 sequence', 'Is required; hidden prereq for Math sequence', 'Is required; hidden prereq for Math sequence in 2nd degree', 'Is required; hidden prereq for required Science sequence', 'Is required; Spring 2022 is 2nd attempt', 'Is required; will replace transfer grade', 'Is required; Spring 2022 is 5th attempt; can\'t pay', 'Required for Degree, Substituiton incoming from Dean', 'Is Required; Repeat being used to boost GPA to minimum 2.0', 'Is required, needed to meet residency', 'Is required, DW has been updated', 'Is required; DW is showing correctly', 'Is required, student dropped class', 'Is required, DW is being corrected', 'Is required, DW is showing correctly, 2ND MAJOR CORRECTED', '1st 7 week, was required, student failed course', '1st 7 weeks, was required for degree', 'Required; 7 week course-student failed']}

    sheet2 = pd.DataFrame(data=comments)

    # Write DataFrame to Excel File
    writer = pd.ExcelWriter(infile+'.xlsx')

    with writer as writer:
        wksht.to_excel(writer, sheet_name='CPoS Report', index=False)
        sheet2.to_excel(writer, sheet_name='Standard Comments', index=False)

    f = infile+'.xlsx'
    wb = opxl.load_workbook(f, read_only=False)

    sheets = wb.worksheets
    for sheet in sheets:
        sheet.column_dimensions[get_column_letter(1)].width = 100
        sheet.column_dimensions[get_column_letter(2)].width = 70

    wb.close()

    # Give confirmation and let the user know the file is ready to be worked
    path = os.path.join(cwd, termcode, nextdir2)

    print(
        f'\nThe file {infile} has now been saved in the {path} directory and is ready to be worked.')
    print('\nHave a great day!\n')

    return False


# Start the Program
while True:
    print('\nWelcome! Enter the needed term code:  ')
    termcode = input()
    question = input(
        '\nDo you need to import multiple files? (\'y\' or \'n\'):  ').lower()
    nextdir2 = input(
        'Enter the month of the folder in which you want the merged file to be saved, (usually the current month):  ').capitalize()

    if question == 'n':
        main(termcode, nextdir2)
        break

    elif question == 'y':
        path = os.getcwd()
        path = os.path.join(path, termcode, nextdir2)
        file_list = []
        files = os.listdir(path)
        for file in files:
            if file.endswith('.csv'):
                file_list.append(file)
            else:
                pass
        print(f'\nThe files have been imported into this program, next I will convert the files into a single Excel File from CSV.')
        input('\nPress enter to continue: ')
        createMergedMultipleCampsFile(file_list, termcode, nextdir2)

    else:
        print('\nPlease enter "y" for multiple files or "n" if you have a single file to work:  ')
        input('Press enter to restart:  ')
